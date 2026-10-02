from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[3]
OUTPUT_PATH = PROJECT_ROOT / "docs" / "architecture" / "finsight-database-design.pptx"

PALETTE = {
    "background": "F2F6F4",
    "ink": "17313C",
    "muted": "5E737D",
    "line": "C7D4D8",
    "white": "FFFFFF",
    "blue": "3767A6",
    "blue_fill": "EAF1FB",
    "teal": "087E78",
    "teal_fill": "E4F4EF",
    "amber": "C47425",
    "amber_fill": "FFF3DF",
    "rose": "BC594D",
    "rose_fill": "FAECE9",
    "green": "54804C",
    "green_fill": "EDF5EC",
    "violet": "7156A5",
    "violet_fill": "F0ECF8",
    "dark": "183842",
}


def rgb(name):
    return RGBColor.from_string(PALETTE.get(name, name))


def add_text(
    slide,
    x,
    y,
    width,
    height,
    text,
    size=12,
    color="ink",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0.04,
    font="Aptos",
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign

    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        paragraph.space_before = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def add_card(
    slide,
    x,
    y,
    width,
    height,
    title,
    lines,
    accent,
    fill="white",
    title_size=15,
    body_size=10,
    footer=None,
):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = rgb(fill)
    card.line.color.rgb = rgb("line")
    card.line.width = Pt(1)
    card.adjustments[0] = 0.08

    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(0.08),
        Inches(height),
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = rgb(accent)
    strip.line.fill.background()

    add_text(
        slide,
        x + 0.19,
        y + 0.08,
        width - 0.3,
        0.34,
        title,
        size=title_size,
        color="ink",
        bold=True,
    )
    body_height = height - (0.76 if footer else 0.5)
    add_text(
        slide,
        x + 0.19,
        y + 0.47,
        width - 0.3,
        body_height,
        "\n".join(lines),
        size=body_size,
        color="muted",
        valign=MSO_ANCHOR.TOP,
        margin=0.02,
    )
    if footer:
        add_text(
            slide,
            x + 0.19,
            y + height - 0.28,
            width - 0.3,
            0.19,
            footer,
            size=8,
            color=accent,
            bold=True,
            margin=0.01,
        )
    return card


def add_line(slide, x1, y1, x2, y2, line_color="line", width=1.4):
    line = slide.shapes.add_connector(
        1,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(line_color)
    line.line.width = Pt(width)
    return line


def add_arrow(slide, x, y, accent):
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(x),
        Inches(y),
        Inches(0.2),
        Inches(0.14),
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = rgb(accent)
    arrow.line.fill.background()


def add_header(slide, number, title, subtitle):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("background")
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(0.33),
        Inches(0.08),
        Inches(0.83),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = rgb("teal")
    accent.line.fill.background()
    add_text(
        slide,
        0.7,
        0.31,
        11.7,
        0.2,
        f"FINSIGHT AI  /  DATABASE DESIGN  /  {number}",
        size=9,
        color="teal",
        bold=True,
    )
    add_text(
        slide,
        0.7,
        0.54,
        12.0,
        0.45,
        title,
        size=26,
        color="ink",
        bold=True,
        font="Aptos Display",
    )
    add_text(
        slide,
        0.7,
        1.02,
        12.0,
        0.28,
        subtitle,
        size=11,
        color="muted",
    )
    add_text(
        slide,
        12.45,
        7.14,
        0.35,
        0.18,
        number,
        size=9,
        color="muted",
        align=PP_ALIGN.RIGHT,
    )


def add_footer_band(slide, y, text, size=12):
    band = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.55),
        Inches(y),
        Inches(12.25),
        Inches(0.72),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = rgb("dark")
    band.line.fill.background()
    add_text(
        slide,
        0.78,
        y + 0.08,
        11.8,
        0.55,
        text,
        size=size,
        color="white",
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def slide_domain_overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "01",
        "Four database domains, one traceable workflow",
        "Source financial facts feed investigations; AI and human-review data stay linked to the case.",
    )

    add_line(slide, 3.35, 2.5, 4.05, 2.5, "blue")
    add_text(slide, 3.32, 2.22, 0.78, 0.2, "user / role", size=8, color="blue", bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 8.9, 2.5, 9.65, 2.5, "amber")
    add_text(slide, 8.9, 2.22, 0.75, 0.2, "settlement FK", size=8, color="amber", bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 8.9, 5.15, 9.65, 5.15, "rose")
    add_text(slide, 8.91, 4.87, 0.74, 0.2, "case / run", size=8, color="rose", bold=True, align=PP_ALIGN.CENTER)

    add_card(
        slide,
        0.55,
        1.55,
        2.8,
        2.1,
        "AUTHENTICATION",
        ["users", "roles", "user_role_mappings", "permissions"],
        "blue",
        fill="blue_fill",
        body_size=11,
        footer="Identity and assigned roles",
    )
    add_card(
        slide,
        4.05,
        1.55,
        4.85,
        3.85,
        "FINANCIAL SOURCE DATA",
        [
            "customers -> accounts",
            "accounts -> portfolios -> holdings",
            "securities -> orders / holdings",
            "orders -> executions + transactions",
            "execution_transactions (allocation bridge)",
            "transactions -> settlements -> cash_movements",
        ],
        "amber",
        fill="amber_fill",
        body_size=11,
        footer="System-of-record financial facts",
    )
    add_card(
        slide,
        9.65,
        1.55,
        3.1,
        2.1,
        "INVESTIGATION",
        ["cases", "findings", "evidence", "approval_events"],
        "rose",
        fill="rose_fill",
        body_size=11,
        footer="Discrepancy and review lifecycle",
    )
    add_card(
        slide,
        9.65,
        4.0,
        3.1,
        2.1,
        "AI GOVERNANCE",
        ["policy_documents -> policy_chunks", "explanation_runs"],
        "violet",
        fill="violet_fill",
        body_size=11,
        footer="Grounding and provenance",
    )
    add_text(
        slide,
        0.68,
        6.5,
        12.0,
        0.32,
        "22 live tables across authentication, financial, investigation, and ai schemas.",
        size=12,
        color="ink",
        bold=True,
    )


def slide_financial_chain(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "02",
        "Financial records follow the transaction lifecycle",
        "The reconciliation compares a settlement amount with the completed cash movements linked to it.",
    )

    items = [
        (0.42, "CUSTOMERS", "customer_id PK"),
        (2.55, "ACCOUNTS", "account_id PK\ncustomer_id FK"),
        (4.68, "ORDERS", "order_id PK\naccount_id / security_id FK"),
        (6.81, "TRANSACTIONS", "transaction_id PK\norder_id FK"),
        (8.94, "SETTLEMENTS", "settlement_id PK\ntransaction_id FK"),
        (11.07, "CASH MOVEMENTS", "cash_movement_id PK\nsettlement_id / account_id FK"),
    ]
    y = 2.05
    width = 1.84
    height = 1.18

    for x, title, description in items:
        add_card(
            slide,
            x,
            y,
            width,
            height,
            title,
            description.split("\n"),
            "amber",
            title_size=10,
            body_size=8,
        )

    for index in range(len(items) - 1):
        x1 = items[index][0] + width
        x2 = items[index + 1][0]
        add_line(slide, x1 + 0.02, y + 0.58, x2 - 0.04, y + 0.58, "amber", 1.5)
        add_arrow(slide, x2 - 0.19, y + 0.51, "amber")

    support = [
        (0.55, "PORTFOLIOS", ["account_id FK", "portfolio_id PK"], "blue", "blue_fill"),
        (3.0, "HOLDINGS", ["portfolio_id FK", "security_id FK"], "blue", "blue_fill"),
        (5.45, "SECURITIES", ["security_id PK", "used by orders / holdings"], "blue", "blue_fill"),
        (7.9, "EXECUTIONS", ["execution_id PK", "order_id FK"], "teal", "teal_fill"),
        (10.35, "EXECUTION_TRANSACTIONS", ["composite PK", "allocation bridge"], "violet", "violet_fill"),
    ]
    support_y = 3.85
    support_w = 2.15
    for x, title, lines, accent, fill in support:
        add_card(
            slide,
            x,
            support_y,
            support_w,
            1.15,
            title,
            lines,
            accent,
            fill=fill,
            title_size=10,
            body_size=8,
        )

    add_text(slide, 2.5, 5.12, 2.9, 0.22, "account_id links portfolio", size=8, color="blue", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 5.1, 5.12, 2.7, 0.22, "security_id links holdings / orders", size=8, color="blue", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 8.9, 5.12, 3.0, 0.22, "execution + transaction allocation", size=8, color="violet", bold=True, align=PP_ALIGN.CENTER)

    add_footer_band(
        slide,
        5.75,
        "RECONCILIATION: settlements.expected_cash_amount  vs.  signed sum of completed cash_movements  ->  MATCH / MISMATCH",
        size=11,
    )
    add_text(
        slide,
        0.62,
        6.65,
        12.0,
        0.24,
        "Relationships shown with arrows are declared foreign keys; execution_transactions has a composite key over execution_id + transaction_id.",
        size=9,
        color="muted",
    )


def slide_investigation_ai(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "03",
        "Investigation and AI data are anchored to the case",
        "Each explanation and approval can be traced to its source records, retrieved policy, actor, and outcome.",
    )

    add_card(
        slide,
        0.55,
        1.65,
        2.65,
        1.15,
        "FINANCIAL SETTLEMENT",
        ["settlement_id PK", "expected / settled cash"],
        "amber",
        fill="amber_fill",
        title_size=12,
        body_size=9,
    )
    add_card(
        slide,
        3.65,
        2.55,
        2.65,
        1.4,
        "INVESTIGATION CASE",
        ["investigation_id PK", "settlement_id FK", "status / discrepancy"],
        "rose",
        fill="rose_fill",
        title_size=13,
        body_size=9,
    )
    add_card(
        slide,
        0.55,
        4.0,
        2.65,
        1.15,
        "FINDINGS + EVIDENCE",
        ["cause / severity / action", "source + evidence text"],
        "rose",
        title_size=12,
        body_size=9,
    )
    add_card(
        slide,
        7.05,
        1.65,
        2.65,
        1.25,
        "AI EXPLANATION RUN",
        ["run_id PK; provider / model", "input snapshot / hash / output"],
        "violet",
        fill="violet_fill",
        title_size=12,
        body_size=9,
    )
    add_card(
        slide,
        10.1,
        1.65,
        2.65,
        1.25,
        "POLICY DOCUMENTS",
        ["approved version / dates", "approved_by_user_id FK"],
        "teal",
        fill="teal_fill",
        title_size=12,
        body_size=9,
    )
    add_card(
        slide,
        10.1,
        3.3,
        2.65,
        1.15,
        "POLICY CHUNKS",
        ["document_id FK", "searchable text + keywords"],
        "teal",
        fill="teal_fill",
        title_size=12,
        body_size=9,
    )
    add_card(
        slide,
        7.05,
        4.7,
        2.65,
        1.15,
        "APPROVAL EVENTS",
        ["actor / comment / status transition", "optional explanation_run_id FK"],
        "green",
        fill="green_fill",
        title_size=12,
        body_size=8,
    )

    add_line(slide, 3.2, 2.23, 3.65, 2.85, "amber")
    add_arrow(slide, 3.47, 2.58, "amber")
    add_line(slide, 3.2, 4.55, 3.65, 3.6, "rose")
    add_arrow(slide, 3.47, 3.72, "rose")
    add_line(slide, 6.3, 2.88, 7.05, 2.3, "violet")
    add_arrow(slide, 6.82, 2.45, "violet")
    add_line(slide, 8.38, 2.9, 8.38, 4.7, "green")
    add_arrow(slide, 8.31, 4.48, "green")
    add_line(slide, 10.1, 2.28, 9.7, 2.28, "teal")
    add_line(slide, 11.42, 2.9, 11.42, 3.3, "teal")
    add_arrow(slide, 11.35, 3.12, "teal")
    add_text(slide, 9.6, 2.96, 1.5, 0.19, "retrieved context", size=8, color="teal", bold=True, align=PP_ALIGN.CENTER)

    add_footer_band(
        slide,
        6.15,
        "CASE FLOW: OPEN / INVESTIGATING -> PENDING_APPROVAL -> APPROVED or REJECTED; REQUEST_CHANGES returns to INVESTIGATING",
        size=10,
    )
    add_text(
        slide,
        0.63,
        6.94,
        12.0,
        0.17,
        "Audit runs and approval events are append-only. Evidence source_table/source_record_id is a polymorphic reference, not an enforced FK.",
        size=8,
        color="muted",
    )


def build_deck():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "FinSight AI - Database Design"
    presentation.core_properties.subject = "Database schemas, relationships, and workflow"
    presentation.core_properties.author = "FinSight AI"

    slide_domain_overview(presentation)
    slide_financial_chain(presentation)
    slide_investigation_ai(presentation)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_deck()
